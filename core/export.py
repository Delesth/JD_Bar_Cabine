"""Export Excel des tableaux de bord sur une période."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.calculs import (COMPTES, OPERATEURS, TYPES_CABINE, Donnees, indicateurs,
                          indicateurs_cabine, par_produit, rejouer, soldes, total_mouvements)
from core.format import SIGNATURE

POLICE = "Arial"
VERT = "1E5A42"
FCFA = '#,##0;-#,##0;"-"'
QTE = '#,##0.##;-#,##0.##;"-"'
PCT = '0.0%;-0.0%;"-"'
DATE = "DD/MM/YYYY"
TYPES_VENTE = {"normal": "Prix normal", "reduction": "Tarif réduit", "perte": "Vente à perte"}
STATUTS = {"validee": "Validée", "en_attente": "En attente", "refusee": "Refusée"}
TYPES_MVT = {"apport": "Apport", "depense": "Dépense", "decaissement": "Décaissement",
             "reevaluation": "Réévaluation", "transfert": "Transfert interne"}


def _dans(x, debut, fin):
    return (debut is None or x >= debut) and (fin is None or x <= fin)


def _libelle_periode(debut, fin):
    if debut is None:
        return "Depuis le début"
    return f"Du {debut.strftime('%d/%m/%Y')} au {fin.strftime('%d/%m/%Y')}"


def _entete(ws, titre, periode):
    ws["A1"] = titre
    ws["A1"].font = Font(name=POLICE, bold=True, size=14, color=VERT)
    ws["A2"] = periode
    ws["A2"].font = Font(name=POLICE, italic=True, size=10, color="555555")


def _tableau(ws, ligne, colonnes, lignes, totaux=()):
    """Écrit un tableau à partir de la ligne donnée. colonnes : [(titre, format, largeur)].
    totaux : titres des colonnes à totaliser par une formule SOMME. Renvoie la ligne suivante."""
    fin_bord = Side(style="thin", color="D5DDD7")
    for j, (titre, _, largeur) in enumerate(colonnes, start=1):
        c = ws.cell(row=ligne, column=j, value=titre)
        c.font = Font(name=POLICE, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERT)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        lettre = get_column_letter(j)
        ws.column_dimensions[lettre].width = max(ws.column_dimensions[lettre].width or 0, largeur)
    debut_donnees = ligne + 1
    for i, valeurs in enumerate(lignes, start=debut_donnees):
        for j, ((_, fmt, _), v) in enumerate(zip(colonnes, valeurs), start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=POLICE)
            c.border = Border(bottom=fin_bord)
            if fmt:
                c.number_format = fmt
    suivante = debut_donnees + len(lignes)
    if totaux and lignes:
        ws.cell(row=suivante, column=1, value="Total").font = Font(name=POLICE, bold=True)
        for j, (titre, fmt, _) in enumerate(colonnes, start=1):
            if titre in totaux:
                lettre = get_column_letter(j)
                c = ws.cell(row=suivante, column=j,
                            value=f"=SUM({lettre}{debut_donnees}:{lettre}{suivante - 1})")
                c.font = Font(name=POLICE, bold=True)
                c.number_format = fmt or FCFA
        suivante += 1
    if not lignes:
        ws.cell(row=suivante, column=1, value="Aucune donnée sur la période").font = Font(
            name=POLICE, italic=True, color="888888")
        suivante += 1
    ws.freeze_panes = ws.cell(row=debut_donnees, column=1)
    return suivante


def excel_tableau_de_bord(d: Donnees, debut: date | None, fin: date | None, proprietaire: bool) -> bytes:
    r = rejouer(d)
    ib = indicateurs(d, r, debut, fin)
    ic = indicateurs_cabine(d, debut, fin)
    dep_bar = total_mouvements(d, "depense", debut, fin, source="caisse_bar")
    periode = _libelle_periode(debut, fin)
    nomp = {p.id: p.nom for p in d.produits}
    unite = {p.id: p.unite_vente for p in d.produits}
    wb = Workbook()

    # ---- Synthèse
    ws = wb.active
    ws.title = "Synthèse"
    _entete(ws, "Bar & Cabine — synthèse", periode)
    lignes = [
        ("Bar", "Chiffre d'affaires", ib.ca, FCFA),
        ("Bar", "Coût des boissons vendues", ib.cout_ventes, FCFA),
        ("Bar", "Pertes de stock (comptages)", -ib.pertes_stock, FCFA),
        ("Bar", "Dépenses du bar", dep_bar, FCFA),
        ("Bar", "Bénéfice net du bar", ib.benefice - dep_bar, FCFA),
        ("Bar", "Marge sur ventes", ib.marge / 100, PCT),
        ("Bar", "Achats de la période", ib.achats, FCFA),
        ("Bar", "Réductions accordées", ib.reductions, FCFA),
        ("Bar", "Part des réductions dans le CA au prix normal", ib.part_reductions / 100, PCT),
    ]
    for k, nom in OPERATEURS.items():
        lignes += [("Cabine", f"Commissions {nom}", ic.commissions[k], FCFA),
                   ("Cabine", f"Volume d'opérations {nom}", ic.volumes[k], FCFA)]
    lignes += [
        ("Cabine", "Dépenses de la cabine", ic.depenses, FCFA),
        ("Cabine", "Bénéfice de la cabine", ic.benefice, FCFA),
        ("Total", "Chiffre d'affaires total (ventes + commissions)", ib.ca + ic.commission_totale, FCFA),
        ("Total", "Bénéfice total", ib.benefice - dep_bar + ic.benefice, FCFA),
    ]
    if proprietaire:
        sol = soldes(d)
        valeur_stock = sum(max(s["qte"], 0) * s["cmp"] for s in r.stock.values())
        lignes += [("Trésorerie (à ce jour)", nom, sol[k], FCFA) for k, nom in COMPTES.items()]
        lignes += [("Trésorerie (à ce jour)", "Valeur du stock au coût d'achat", valeur_stock, FCFA),
                   ("Période", "Décaissements", total_mouvements(d, "decaissement", debut, fin), FCFA)]
    ligne = 4
    for j, (t, l) in enumerate([("Activité", 26), ("Indicateur", 48), ("Valeur", 18)], start=1):
        c = ws.cell(row=ligne, column=j, value=t)
        c.font = Font(name=POLICE, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=VERT)
        ws.column_dimensions[get_column_letter(j)].width = l
    for i, (act, lib, val, fmt) in enumerate(lignes, start=ligne + 1):
        ws.cell(row=i, column=1, value=act).font = Font(name=POLICE, bold=True)
        ws.cell(row=i, column=2, value=lib).font = Font(name=POLICE)
        c = ws.cell(row=i, column=3, value=round(val, 4) if fmt == PCT else round(val))
        c.font, c.number_format = Font(name=POLICE), fmt
    bas = ligne + len(lignes) + 2
    ws.cell(row=bas, column=1, value="Montants en FCFA, calculés par l'application (stock au coût moyen "
                                      "pondéré). Le CA de la cabine correspond aux commissions.").font = Font(
        name=POLICE, italic=True, size=9, color="555555")
    ws.cell(row=bas + 1, column=1, value=f"Exporté le {date.today().strftime('%d/%m/%Y')} · {SIGNATURE}").font = Font(
        name=POLICE, italic=True, size=9, color="555555")
    ws.freeze_panes = "A5"

    # ---- Par produit
    ws = wb.create_sheet("Par produit")
    _entete(ws, "Ventes et achats par produit", periode)
    cols = [("Produit", None, 22), ("Quantité vendue", QTE, 12), ("Unité", None, 11),
            ("Chiffre d'affaires (FCFA)", FCFA, 16), ("Bénéfice (FCFA)", FCFA, 15),
            ("Réductions (FCFA)", FCFA, 14), ("Quantité achetée", QTE, 12), ("Achats (FCFA)", FCFA, 15),
            ("Stock actuel", QTE, 11)]
    if proprietaire:
        cols += [("Coût moyen (FCFA)", FCFA, 13), ("Valeur du stock (FCFA)", FCFA, 15)]
    rows = []
    for l in par_produit(d, r, debut, fin):
        p = l["produit"]
        row = [p.nom, l["qte_vendue"], p.unite_vente + "s", round(l["ca"]), round(l["benefice"]),
               round(l["reductions"]), l["qte_achetee"], round(l["achats"]), l["stock"]]
        if proprietaire:
            row += [round(l["cmp"]), round(l["valeur_stock"])]
        rows.append(row)
    _tableau(ws, 4, cols, rows, totaux=("Chiffre d'affaires (FCFA)", "Bénéfice (FCFA)", "Réductions (FCFA)",
                                        "Achats (FCFA)", "Valeur du stock (FCFA)"))

    # ---- Ventes
    ws = wb.create_sheet("Ventes")
    _entete(ws, "Ventes du bar", periode)
    ventes = sorted([v for v in d.ventes if _dans(v.date, debut, fin)], key=lambda v: (v.date, v.id))
    cols = [("Date", DATE, 12), ("Produit", None, 20), ("Type", None, 14), ("Quantité", QTE, 10),
            ("Unité", None, 11), ("Prix normal (FCFA)", FCFA, 14), ("Encaissé (FCFA)", FCFA, 14),
            ("Réduction (FCFA)", FCFA, 13), ("Motif", None, 26), ("Statut", None, 12), ("Saisi par", None, 20)]
    if proprietaire:
        cols.insert(8, ("Bénéfice (FCFA)", FCFA, 13))
    rows = []
    for v in ventes:
        row = [v.date, nomp.get(v.produit_id, "?"), TYPES_VENTE.get(v.type, v.type), v.quantite,
               unite.get(v.produit_id, "") + "s", round(v.montant_normal), round(v.montant_encaisse),
               round(v.reduction)]
        if proprietaire:
            row.append(round(v.montant_encaisse - r.cout_vente.get(v.id, 0)))
        row += [v.motif or "", STATUTS.get(v.statut, v.statut), d.utilisateurs.get(v.auteur_id, "?")]
        rows.append(row)
    _tableau(ws, 4, cols, rows, totaux=("Prix normal (FCFA)", "Encaissé (FCFA)", "Réduction (FCFA)",
                                        "Bénéfice (FCFA)"))

    # ---- Achats
    ws = wb.create_sheet("Achats")
    _entete(ws, "Achats de boissons", periode)
    achats = sorted([a for a in d.achats if _dans(a.date, debut, fin)], key=lambda a: (a.date, a.id))
    cols = [("Date", DATE, 12), ("Produit", None, 20), ("Conditionnement", None, 15), ("Nombre", QTE, 9),
            ("Unités par cond.", QTE, 10), ("Unités ajoutées", QTE, 10), ("Prix d'achat / cond. (FCFA)", FCFA, 15),
            ("Montant (FCFA)", FCFA, 14), ("Coût unitaire (FCFA)", FCFA, 13), ("Prix de vente (FCFA)", FCFA, 13),
            ("Saisi par", None, 20)]
    rows = [[a.date, nomp.get(a.produit_id, "?"), a.conditionnement, a.nb_cond, a.unites_par_cond, a.unites,
             round(a.prix_achat_cond), round(a.montant), round(a.cout_unitaire), round(a.prix_vente_unite),
             d.utilisateurs.get(a.auteur_id, "?")] for a in achats]
    _tableau(ws, 4, cols, rows, totaux=("Montant (FCFA)",))

    # ---- Cabine
    ws = wb.create_sheet("Cabine")
    _entete(ws, "Cabine Airtel / MTN", periode)
    ops = sorted([o for o in d.operations if _dans(o.date, debut, fin)], key=lambda o: (o.date, o.operateur, o.type))
    cols = [("Date", DATE, 12), ("Réseau", None, 10), ("Opération", None, 18), ("Montant (FCFA)", FCFA, 15),
            ("Commission (FCFA)", FCFA, 15), ("Saisi par", None, 20)]
    rows = [[o.date, OPERATEURS.get(o.operateur, o.operateur), TYPES_CABINE.get(o.type, o.type),
             round(o.montant), round(o.commission), d.utilisateurs.get(o.auteur_id, "?")] for o in ops]
    _tableau(ws, 4, cols, rows, totaux=("Montant (FCFA)", "Commission (FCFA)"))

    # ---- Dépenses
    ws = wb.create_sheet("Dépenses")
    _entete(ws, "Dépenses", periode)
    deps = sorted([m for m in d.mouvements if m.type == "depense" and _dans(m.date, debut, fin)],
                  key=lambda m: (m.date, m.id))
    cols = [("Date", DATE, 12), ("Activité", None, 10), ("Catégorie", None, 20), ("Détail", None, 30),
            ("Montant (FCFA)", FCFA, 14), ("Saisi par", None, 20)]
    rows = [[m.date, "Bar" if m.compte_source == "caisse_bar" else "Cabine", m.categorie or "", m.libelle or "",
             round(m.montant), d.utilisateurs.get(m.auteur_id, "?")] for m in deps]
    _tableau(ws, 4, cols, rows, totaux=("Montant (FCFA)",))

    # ---- Trésorerie (propriétaire)
    if proprietaire:
        ws = wb.create_sheet("Trésorerie")
        _entete(ws, "Mouvements de trésorerie", periode)
        mv = sorted([m for m in d.mouvements if _dans(m.date, debut, fin)], key=lambda m: (m.date, m.id))
        cols = [("Date", DATE, 12), ("Type", None, 18), ("De", None, 16), ("Vers", None, 16),
                ("Montant (FCFA)", FCFA, 14), ("Catégorie", None, 16), ("Détail", None, 28), ("Saisi par", None, 20)]
        rows = [[m.date, TYPES_MVT.get(m.type, m.type), COMPTES.get(m.compte_source, ""),
                 COMPTES.get(m.compte_dest, ""), round(m.montant), m.categorie or "", m.libelle or "",
                 d.utilisateurs.get(m.auteur_id, "?")] for m in mv]
        _tableau(ws, 4, cols, rows)

    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def nom_fichier(debut, fin) -> str:
    if debut is None:
        return f"bar-cabine_complet_{date.today().isoformat()}.xlsx"
    return f"bar-cabine_{debut.isoformat()}_au_{fin.isoformat()}.xlsx"
